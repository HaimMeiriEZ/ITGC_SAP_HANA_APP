"""Tests for compensating-controls repository, service, and working-paper sheet."""
from __future__ import annotations

import unittest
from unittest import mock
from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import load_workbook

from src.persistence.compensating_control_repository import CompensatingControlRepository
from src.reporting.working_paper_report import write_control_working_paper
from src.services.compensating_control_service import (
    DEFAULT_COMPENSATING_COLUMN_WIDTHS,
    DEFAULT_COMPENSATING_ROW_HEIGHT,
    build_compensating_control_rows,
    build_findings_brief_summary,
    build_findings_description,
    control_has_findings,
)


class CompensatingControlServiceTests(unittest.TestCase):
    def test_control_has_findings_includes_review_status(self) -> None:
        self.assertTrue(
            control_has_findings([{"status": "לסקירה", "full_description": "משתמש לסקירה"}])
        )

    def test_control_has_findings_excludes_passing_status(self) -> None:
        self.assertFalse(
            control_has_findings([{"status": "תקין", "full_description": "תקין"}])
        )

    def test_build_findings_description_joins_non_passing_rows(self) -> None:
        text = build_findings_description(
            [
                {"status": "עם ממצא", "full_description": "ממצא א"},
                {"status": "לסקירה", "full_description": "ממצא ב"},
                {"status": "תקין", "full_description": "לא אמור להיכלל"},
            ]
        )
        self.assertIn("ממצא א", text)
        self.assertIn("ממצא ב", text)
        self.assertNotIn("לא אמור להיכלל", text)

    def test_build_findings_brief_summary_for_review_users(self) -> None:
        text = build_findings_brief_summary(
            [
                {"status": "לסקירה", "user_name": "USER1"},
                {"status": "לסקירה", "user_name": "USER2"},
            ],
            control_id="MA7-17_AYALON_30",
            control_meta={"check_type": "סקירת הרשאות משתמשים"},
        )
        self.assertIn("נמצאו", text)
        self.assertIn("2", text)
        self.assertIn("סקיר", text)

    def test_build_findings_brief_summary_for_permissions(self) -> None:
        text = build_findings_brief_summary(
            [
                {"status": "עם ממצא", "user_name": "A"},
                {"status": "עם ממצא", "user_name": "B"},
                {"status": "עם ממצא", "user_name": "C"},
            ],
            control_meta={"check_type": "הרשאות ניהול משתמשים"},
        )
        self.assertEqual(text, "נמצאו 3 משתמשים בעלי הרשאות רגישות.")

    def test_build_compensating_control_rows_filters_in_scope_with_findings(self) -> None:
        summary = {
            "MA7-17_AYALON_30": {"control_id": "MA7-17_AYALON_30"},
            "MA2-2_AYALON_6": {"control_id": "MA2-2_AYALON_6"},
        }
        details = {
            "MA7-17_AYALON_30": [{"status": "לסקירה", "full_description": "לסקירה"}],
            "MA2-2_AYALON_6": [{"status": "תקין", "full_description": "תקין"}],
        }

        def _meta(control_id: str) -> dict[str, str]:
            return {
                "description": f"desc-{control_id}",
                "risk_description": f"risk-{control_id}",
            }

        rows = build_compensating_control_rows(
            summary,
            details,
            {},
            _meta,
            lambda cid: cid == "MA7-17_AYALON_30",
        )
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["control_id"], "MA7-17_AYALON_30")
        self.assertIn("נמצאו", rows[0]["findings_brief"])


class CompensatingControlRepositoryTests(unittest.TestCase):
    def test_attach_replaces_previous_file(self) -> None:
        with TemporaryDirectory() as temp_dir:
            base = Path(temp_dir)
            output_dir = base / "output"
            output_dir.mkdir()
            repo = CompensatingControlRepository(output_dir, base)
            data: dict[str, dict] = {}

            first_source = base / "first.pdf"
            second_source = base / "second.pdf"
            first_source.write_text("first", encoding="utf-8")
            second_source.write_text("second", encoding="utf-8")

            first_entry = repo.attach_file("MA7-17_AYALON_30", first_source, data)
            first_stored = Path(first_entry["stored_path"])
            self.assertTrue(first_stored.exists())

            second_entry = repo.attach_file("MA7-17_AYALON_30", second_source, data)
            second_stored = Path(second_entry["stored_path"])
            self.assertTrue(second_stored.exists())
            self.assertFalse(first_stored.exists())
            self.assertEqual(data["MA7-17_AYALON_30"]["original_filename"], "second.pdf")


class CompensatingControlWorkingPaperTests(unittest.TestCase):
    _SUMMARY = {
        "control_id": "MA7-17_AYALON_30",
        "source_file": "USR02",
        "extraction_date": "2026-07-22",
        "total_records": 1,
        "finding_records": 1,
        "description": "סקירת הרשאות",
    }
    _DETAIL = [
        {
            "control_id": "MA7-17_AYALON_30",
            "status": "לסקירה",
            "full_description": "משתמש לסקירה",
        }
    ]

    def test_working_paper_without_entry_has_no_compensating_sheet(self) -> None:
        with TemporaryDirectory() as temp_dir:
            output_path = Path(temp_dir) / "wp.xlsx"
            write_control_working_paper(
                control_id="MA7-17_AYALON_30",
                summary_record=self._SUMMARY,
                detail_rows=self._DETAIL,
                raw_population_rows=[],
                ipe_entries=[],
                work_environment_label="ייצור",
                output_path=output_path,
            )
            workbook = load_workbook(output_path)
            self.assertNotIn("בקרה מפצה", workbook.sheetnames)

    def test_working_paper_with_image_entry_embeds_compensating_sheet(self) -> None:
        with TemporaryDirectory() as temp_dir:
            base = Path(temp_dir)
            image_path = base / "evidence.png"
            image_path.write_bytes(
                b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01"
                b"\x00\x00\x00\x01\x08\x02\x00\x00\x00\x90wS\xde\x00\x00\x00\x0cIDATx\x9cc\xf8\x0f\x00\x00\x01\x01\x00\x05\x18\xd8N\x00\x00\x00\x00IEND\xaeB`\x82"
            )
            output_path = base / "wp.xlsx"
            write_control_working_paper(
                control_id="MA7-17_AYALON_30",
                summary_record=self._SUMMARY,
                detail_rows=self._DETAIL,
                raw_population_rows=[],
                ipe_entries=[],
                work_environment_label="ייצור",
                output_path=output_path,
                compensating_control_entry={
                    "control_id": "MA7-17_AYALON_30",
                    "original_filename": "evidence.png",
                    "stored_path": str(image_path),
                    "added_at": "2026-07-22T10:00:00",
                },
            )
            workbook = load_workbook(output_path)
            self.assertIn("בקרה מפצה", workbook.sheetnames)
            sheet = workbook["בקרה מפצה"]
            self.assertEqual(sheet["B3"].value, "evidence.png")

    def test_working_paper_with_pdf_requests_ole_embed_without_hyperlink(self) -> None:
        with TemporaryDirectory() as temp_dir:
            base = Path(temp_dir)
            pdf_path = base / "evidence.pdf"
            pdf_path.write_text("%PDF-1.4 test", encoding="utf-8")
            output_path = base / "wp.xlsx"
            with mock.patch(
                "src.reporting.working_paper_report.embed_file_in_worksheet",
                return_value=True,
            ) as embed_mock:
                write_control_working_paper(
                    control_id="MA7-17_AYALON_30",
                    summary_record=self._SUMMARY,
                    detail_rows=self._DETAIL,
                    raw_population_rows=[],
                    ipe_entries=[],
                    work_environment_label="ייצור",
                    output_path=output_path,
                    compensating_control_entry={
                        "control_id": "MA7-17_AYALON_30",
                        "original_filename": "evidence.pdf",
                        "stored_path": str(pdf_path),
                        "added_at": "2026-07-22T10:00:00",
                    },
                )
                embed_mock.assert_called_once()
            copied = base / "MA7-17_AYALON_30_compensating_evidence.pdf"
            self.assertFalse(copied.exists())
            workbook = load_workbook(output_path)
            sheet = workbook["בקרה מפצה"]
            self.assertIsNone(sheet["B6"].hyperlink)


if __name__ == "__main__":
    unittest.main()
